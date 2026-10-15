# -*- coding: utf-8 -*-
"""把 docx 中的 LaTeX 公式（$$..$$ 显示式与 $..$ 行内式）转为 Word 原生公式 (OMML)。"""
import re
import copy
import latex2mathml.converter
import docx
from lxml import etree
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.shared import Pt

DOCX = r'E:\project\web-project-workflow\论文\面向AI辅助软件开发的图谱驱动上下文工程方法研究与系统实现-前5章.docx'
XSL = r'C:\Program Files\Microsoft Office\root\Office16\MML2OMML.XSL'

W = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
xslt = etree.XSLT(etree.parse(XSL))

def latex_to_omml(latex):
    mathml = latex2mathml.converter.convert(latex)
    tree = etree.fromstring(mathml)
    omml = xslt(tree)
    root = omml.getroot()
    # 根元素应为 m:oMath；若被包了一层取子元素
    if root.tag.endswith('}oMath') or root.tag.endswith('}oMathPara'):
        return root
    for child in root.iter():
        if child.tag.endswith('}oMath'):
            return child
    raise ValueError('no oMath produced')

d = docx.Document(DOCX)
converted_disp = converted_inl = failed = 0
fail_list = []

def rpr_template(p):
    for r in p.runs:
        el = r._r.find(qn('w:rPr'))
        if el is not None:
            return copy.deepcopy(el)
    return None

for p in d.paragraphs:
    text = p.text
    if '$' not in text:
        continue

    # --- 显示公式：整段是 $$...$$ ---
    m = re.fullmatch(r'\s*\$\$(.+?)\$\$\s*', text, re.S)
    if m:
        latex = m.group(1).strip()
        tag = re.search(r'\\tag\{([^}]+)\}', latex)
        latex_clean = re.sub(r'\s*\\tag\{[^}]+\}', '', latex).strip()
        try:
            omath = latex_to_omml(latex_clean)
        except Exception as e:
            failed += 1
            fail_list.append(('disp', latex_clean[:60], str(e)[:80]))
            continue
        rpr = rpr_template(p)
        for r in list(p.runs):
            r._r.getparent().remove(r._r)
        # 清掉段内其他直接子元素（保留 pPr）
        for child in list(p._p):
            if not child.tag.endswith('}pPr'):
                p._p.remove(child)
        # Use a center tab stop for the formula and a right tab stop for its number.
        section = d.sections[0]
        text_width = section.page_width - section.left_margin - section.right_margin
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.first_line_indent = 0
        p.paragraph_format.left_indent = 0
        p.paragraph_format.right_indent = 0
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.tab_stops.add_tab_stop(int(text_width / 2), WD_TAB_ALIGNMENT.CENTER)
        p.paragraph_format.tab_stops.add_tab_stop(int(text_width), WD_TAB_ALIGNMENT.RIGHT)
        p.add_run('\t')
        p._p.append(omath)
        if tag:
            number_run = p.add_run('\t（' + tag.group(1) + '）')
            number_run.font.name = 'Times New Roman'
            number_run.font.size = Pt(10.5)
        converted_disp += 1
        continue

    # --- 行内公式：段落中夹杂 $..$ ---
    if '$$' in text:
        continue  # 非整段的 $$ 不处理
    segs = re.split(r'(\$[^$]+\$)', text)
    if len(segs) < 2:
        continue
    rpr = rpr_template(p)
    new_els = []
    ok = True
    for seg in segs:
        if not seg:
            continue
        if seg.startswith('$') and seg.endswith('$'):
            try:
                omath = latex_to_omml(seg[1:-1])
                new_els.append(omath)
            except Exception as e:
                ok = False
                fail_list.append(('inl', seg[:50], str(e)[:80]))
                new_els = None
                break
        else:
            r = etree.SubElement(etree.Element(qn('w:r')), qn('w:t'))  # placeholder
            new_els.append(('text', seg))
    if not ok or new_els is None:
        continue
    for child in list(p._p):
        if not child.tag.endswith('}pPr'):
            p._p.remove(child)
    for el in new_els:
        if isinstance(el, tuple):
            run = p.add_run(el[1])
            if rpr is not None:
                run._r.insert(0, copy.deepcopy(rpr))
        else:
            p._p.append(el)
    converted_inl += 1

d.save(DOCX)
print(f'显示公式转换: {converted_disp}, 行内公式段落: {converted_inl}, 失败: {failed}')
for f in fail_list:
    print('  FAIL', f)
